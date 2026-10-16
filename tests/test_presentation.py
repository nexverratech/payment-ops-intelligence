"""The presentation builder: an editable deck with native charts and notes on every slide."""

import importlib.util
import sys
from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx", reason="pip install -r requirements-docs.txt")
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    spec = importlib.util.spec_from_file_location("build_presentation", ROOT / "scripts" / "build_presentation.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod                   # dataclasses look the module up while it loads
    spec.loader.exec_module(mod)
    out = tmp_path_factory.mktemp("deck")
    deck = mod.build(out / "deck.pptx", out / "notes.md")
    return mod, pptx.Presentation(deck), (out / "notes.md").read_text(encoding="utf-8")


def test_widescreen_with_every_slide_titled_and_noted(built):
    mod, prs, _ = built
    assert (prs.slide_width, prs.slide_height) == (mod.W, mod.H)
    assert len(prs.slides) == len(mod.SLIDES) >= 12
    for i, slide in enumerate(prs.slides, 1):
        assert slide.shapes.title is not None and slide.shapes.title.text.strip(), f"slide {i} has no title"
        assert len(slide.notes_slide.notes_text_frame.text) > 100, f"slide {i} has no real speaker notes"


def test_everything_is_editable_no_pictures(built):
    _, prs, _ = built
    charts = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.shape_type != MSO_SHAPE_TYPE.PICTURE, "use native shapes and charts, not images"
            charts += shape.has_chart
    assert charts == 3


def test_coordinates_are_whole_numbers(tmp_path, built):
    """OOXML positions are integers; a fraction makes PowerPoint and other readers reject the file."""
    import re
    import zipfile

    mod, _, _ = built
    deck = mod.build(tmp_path / "deck.pptx", None)
    with zipfile.ZipFile(deck) as z:
        for name in z.namelist():
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                xml = z.read(name).decode("utf-8")
                assert not re.search(r'\b(?:x|y|cx|cy|w|h)="-?\d+\.\d+"', xml), name


def test_shapes_stay_on_the_slide(built):
    _, prs, _ = built
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.LINE:
                continue
            assert shape.left >= 0 and shape.top >= 0, f"slide {i}: {shape.name}"
            assert shape.left + shape.width <= prs.slide_width, f"slide {i}: {shape.name} past the right edge"
            assert shape.top + shape.height <= prs.slide_height, f"slide {i}: {shape.name} past the bottom"


def test_notes_file_matches_the_deck(built):
    _, prs, notes = built
    for i, slide in enumerate(prs.slides, 1):
        assert f"## {i}. {slide.shapes.title.text}" in notes
        assert slide.notes_slide.notes_text_frame.text in notes


def test_committed_notes_are_up_to_date(built):
    _, _, notes = built
    committed = (ROOT / "docs" / "presentation_notes.md").read_text(encoding="utf-8")
    assert committed == notes, "run: python scripts/build_presentation.py"
